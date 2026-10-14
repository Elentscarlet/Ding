#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
打卡记录查询页（钉钉工作台 H5 微应用）+ Excel 导出。

路由：
  GET  /                  查询页（手机/PC 自适应，无需登录即可打开）
  GET  /api/auth          免登所需参数（corpId / clientId）
  POST /api/login         用免登授权码换 userid，校验白名单后下发会话 cookie
  GET  /api/me            当前会话信息
  GET  /api/places        已记录过的地点列表（页面下拉用）
  GET  /api/records       查询：返回汇总 + 明细（JSON）
  GET  /api/export.xlsx   导出 Excel（汇总 + 明细 两个 sheet）

鉴权：
  auth_mode=dingtalk（默认）：钉钉免登拿 userid，必须在 web.admin_user_ids 白名单内；
    会话用 HMAC-SHA256 签名的 cookie 保存，避免每次进页面都走免登。
  auth_mode=none：不做鉴权，仅供本机调试（启动时打醒目告警）。

免登说明：dd.requestAuthCode 属于「无需鉴权即可直接调用」的 JSAPI，
所以本页不需要 dd.config / jsapi_ticket 签名。
"""

import hashlib
import hmac
import io
import json
import logging
import secrets
import threading
import time
from datetime import datetime, timedelta
from datetime import time as dtime
from datetime import timezone
from http.cookies import SimpleCookie
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from typing import List, Optional, Tuple
from urllib.parse import parse_qs, quote, urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SESSION_COOKIE = "ding_session"
DEFAULT_RANGE_DAYS = 7

# 明细/汇总都用的表头
SUMMARY_HEADERS = ["地点", "日期", "首次打卡", "末次打卡", "到岗人数", "到岗人"]
DETAIL_HEADERS = ["姓名", "userId", "地点", "打卡日期", "打卡时间",
                  "距离(米)", "打卡地址", "外勤备注", "打卡结果", "数据来源"]


# ======================================================================
# 纯函数区（可单测）
# ======================================================================
def sign_session(user_id: str, expire_at: float, secret: str) -> str:
    """生成会话 cookie：userId.过期时间戳.HMAC。"""
    if "." in user_id:
        raise ValueError("userId 不能包含 '.'")
    payload = f"{user_id}.{int(expire_at)}"
    signature = hmac.new(secret.encode(), payload.encode(), hashlib.sha256).hexdigest()
    return f"{payload}.{signature}"


def verify_session(token: Optional[str], secret: str,
                   now: Optional[float] = None) -> Optional[str]:
    """校验会话 cookie，通过则返回 userId，否则 None。"""
    if not token:
        return None
    parts = token.split(".")
    if len(parts) != 3:
        return None
    user_id, expire_text, signature = parts
    payload = f"{user_id}.{expire_text}"
    expected = hmac.new(secret.encode(), payload.encode(), hashlib.sha256).hexdigest()
    if not hmac.compare_digest(expected, signature):
        return None
    try:
        expire_at = int(expire_text)
    except ValueError:
        return None
    if (now if now is not None else time.time()) > expire_at:
        return None
    return user_id


def resolve_range(from_text: Optional[str], to_text: Optional[str],
                  tz: timezone,
                  default_days: int = DEFAULT_RANGE_DAYS
                  ) -> Tuple[object, object, datetime, datetime]:
    """把 YYYY-MM-DD 的起止日期换算成 UTC 的 [start, end] 时间区间。

    :return: (起始日期, 结束日期, start_utc, end_utc)
    """
    today = datetime.now(tz).date()
    try:
        to_date = datetime.strptime((to_text or "").strip(), "%Y-%m-%d").date()
    except (TypeError, ValueError):
        to_date = today
    try:
        from_date = datetime.strptime((from_text or "").strip(), "%Y-%m-%d").date()
    except (TypeError, ValueError):
        from_date = to_date - timedelta(days=default_days - 1)

    if from_date > to_date:
        from_date, to_date = to_date, from_date

    start_local = datetime.combine(from_date, dtime.min, tzinfo=tz)
    end_local = datetime.combine(to_date, dtime.max, tzinfo=tz)
    return from_date, to_date, start_local.astimezone(timezone.utc), end_local.astimezone(timezone.utc)


def _to_local(moment, tz: timezone) -> Optional[datetime]:
    if not isinstance(moment, datetime):
        return None
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=timezone.utc)
    return moment.astimezone(tz)


def summarize(records: List[dict], tz: timezone) -> List[dict]:
    """按「地点 + 日期」聚合，得到 地点/日期/首次/末次/人数/到岗人。"""
    buckets: dict = {}
    for record in records:
        local = _to_local(record.get("user_check_time"), tz)
        if local is None:
            continue
        place = record.get("matched_place") or "(未知地点)"
        day = local.strftime("%Y-%m-%d")
        bucket = buckets.setdefault((place, day), {"users": set(), "times": [], "count": 0})
        bucket["count"] += 1
        bucket["times"].append(local)
        user_id = record.get("user_id")
        if user_id:
            bucket["users"].add(user_id)

    rows = []
    for place, day in sorted(buckets, key=lambda item: (item[1], item[0])):
        bucket = buckets[(place, day)]
        rows.append({
            "place": place,
            "date": day,
            "first_time": min(bucket["times"]).strftime("%H:%M:%S"),
            "last_time": max(bucket["times"]).strftime("%H:%M:%S"),
            "user_ids": sorted(bucket["users"]),
            "user_count": len(bucket["users"]),
            "record_count": bucket["count"],
        })
    return rows


def detail_rows(records: List[dict], tz: timezone, name_of) -> List[dict]:
    """明细行（按打卡时间正序），name_of(user_id) -> 姓名。"""
    prepared = []
    for record in records:
        local = _to_local(record.get("user_check_time"), tz)
        if local is None:
            continue
        prepared.append((local, record))
    prepared.sort(key=lambda item: item[0])

    rows = []
    for local, record in prepared:
        user_id = record.get("user_id") or ""
        rows.append({
            "name": name_of(user_id),
            "user_id": user_id,
            "place": record.get("matched_place") or "",
            "date": local.strftime("%Y-%m-%d"),
            "time": local.strftime("%H:%M:%S"),
            "distance": record.get("distance_meters"),
            "address": record.get("user_address") or "",
            "remark": record.get("outside_remark") or "",
            "time_result": record.get("time_result") or "",
            "source_type": record.get("source_type") or "",
        })
    return rows


def build_workbook(summary_rows: List[dict], detail_rows: List[dict],
                   meta: dict) -> bytes:
    """生成 xlsx（汇总 + 明细 两个 sheet），返回字节流。

    :param summary_rows: [{place, date, first_time, last_time, user_count, people}]
                         其中 people 已经是"张三、李四"这样的展示串
    :param detail_rows:  明细行（字段见 DETAIL_HEADERS 顺序）
    :param meta:         查询条件，写进汇总表尾部便于回溯口径
    """
    workbook = Workbook()

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")

    def write_sheet(sheet, headers, rows, widths):
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        if rows:
            sheet.auto_filter.ref = sheet.dimensions
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width

    sheet = workbook.active
    sheet.title = "汇总"
    summary_values = [
        [
            row["place"], row["date"], row["first_time"], row["last_time"],
            row["user_count"], row.get("people", ""),
        ]
        for row in summary_rows
    ]
    write_sheet(sheet, SUMMARY_HEADERS, summary_values, [24, 12, 10, 10, 10, 60])

    detail_sheet = workbook.create_sheet("明细")
    detail_values = [
        [
            row["name"], row["user_id"], row["place"], row["date"], row["time"],
            row["distance"], row["address"], row["remark"],
            row["time_result"], row["source_type"],
        ]
        for row in detail_rows
    ]
    write_sheet(detail_sheet, DETAIL_HEADERS, detail_values,
                [12, 24, 24, 12, 10, 10, 32, 20, 10, 10])

    # 把查询条件写在汇总表的尾部，便于回溯口径
    sheet.append([])
    sheet.append(["查询条件", str(meta.get("range_text") or ""),
                  "地点", str(meta.get("place") or "全部"),
                  "导出时间", str(meta.get("exported_at") or "")])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_row_groups(result: dict) -> Tuple[List[dict], List[dict]]:
    """把查询结果转成写 Excel 需要的两组行（到岗人已经是姓名）。"""
    summary = [
        {
            "place": row["place"],
            "date": row["date"],
            "first_time": row["first_time"],
            "last_time": row["last_time"],
            "user_count": row["user_count"],
            "people": "、".join(row.get("users") or []),
        }
        for row in result["summary"]
    ]
    return summary, result["detail"]


# ======================================================================
# HTTP 服务
# ======================================================================
class WebServer:
    def __init__(self, cfg, api, checkin_store, user_store,
                 utc_offset_hours: int, logger: logging.Logger):
        self.cfg = cfg
        self.api = api
        self.checkin_store = checkin_store
        self.user_store = user_store
        self.logger = logger
        self.tz = timezone(timedelta(hours=utc_offset_hours))
        self._secret = ""
        self._httpd: Optional[ThreadingHTTPServer] = None
        self._thread: Optional[threading.Thread] = None

    # ------------------------------------------------------------------
    # 生命周期
    # ------------------------------------------------------------------
    def start(self):
        if self.cfg.auth_mode == "none":
            self.logger.warning(
                "web.auth_mode=none：打卡查询页**未做任何鉴权**，任何人都能读取考勤数据，请勿在生产环境使用"
            )
        if self.cfg.session_secret:
            self._secret = self.cfg.session_secret
        else:
            self._secret = secrets.token_urlsafe(32)
            self.logger.warning(
                "web.session_secret 未配置，已随机生成；进程重启后所有人需要重新免登"
            )

        self._httpd = ThreadingHTTPServer(
            (self.cfg.host, self.cfg.port), _make_handler(self)
        )
        self._httpd.daemon_threads = True
        self._thread = threading.Thread(
            target=self._httpd.serve_forever, kwargs={"poll_interval": 0.5},
            daemon=True, name="checkin-web",
        )
        self._thread.start()
        self.logger.info(
            "打卡查询页已启动 | http://%s:%d/ | 鉴权=%s | 白名单=%d 人",
            self.cfg.host, self.cfg.port, self.cfg.auth_mode, len(self.cfg.admin_user_ids),
        )

    def stop(self):
        if self._httpd is not None:
            self._httpd.shutdown()
            self._httpd.server_close()
            self._httpd = None
        if self._thread is not None:
            self._thread.join(timeout=5)
            self._thread = None
        self.logger.info("打卡查询页已停止")

    # ------------------------------------------------------------------
    # 会话
    # ------------------------------------------------------------------
    def make_session(self, user_id: str) -> str:
        return sign_session(user_id, time.time() + self.cfg.session_ttl, self._secret)

    def current_user(self, cookie_header: Optional[str]) -> Optional[str]:
        if self.cfg.auth_mode == "none":
            return "(debug)"
        if not cookie_header:
            return None
        jar = SimpleCookie()
        try:
            jar.load(cookie_header)
        except Exception:
            return None
        morsel = jar.get(SESSION_COOKIE)
        return verify_session(morsel.value if morsel else None, self._secret)

    def is_allowed(self, user_id: Optional[str]) -> bool:
        if user_id is None:
            return False
        if self.cfg.auth_mode == "none":
            return True
        return user_id in self.cfg.admin_user_ids

    # ------------------------------------------------------------------
    # 业务
    # ------------------------------------------------------------------
    def name_of_factory(self, user_ids: List[str]):
        """返回 name_of(user_id) -> 姓名，查不到就退回 userId。"""
        try:
            mapping = self.user_store.name_map(user_ids)
        except Exception as e:
            self.logger.warning("读取员工姓名缓存失败，改为显示 userId: %s", e)
            mapping = {}
        return lambda uid: mapping.get(uid) or (uid or "-")

    def query(self, from_text, to_text, place):
        from_date, to_date, start_utc, end_utc = resolve_range(from_text, to_text, self.tz)
        records = self.checkin_store.find_records(
            start_utc, end_utc, place=place or None, limit=self.cfg.export_limit
        )
        summary = summarize(records, self.tz)
        user_ids = sorted({r.get("user_id") for r in records if r.get("user_id")})
        name_of = self.name_of_factory(user_ids)
        detail = detail_rows(records, self.tz, name_of)
        return {
            "from": from_date.strftime("%Y-%m-%d"),
            "to": to_date.strftime("%Y-%m-%d"),
            "place": place or "",
            "record_total": len(records),
            "user_total": len(user_ids),
            "summary": [
                {
                    "place": row["place"],
                    "date": row["date"],
                    "first_time": row["first_time"],
                    "last_time": row["last_time"],
                    "user_count": row["user_count"],
                    "record_count": row["record_count"],
                    "users": [name_of(uid) for uid in row["user_ids"]],
                }
                for row in summary
            ],
            "detail": detail,
        }


def _make_handler(app: WebServer):
    """构造绑定了应用依赖的请求处理器。"""

    class Handler(BaseHTTPRequestHandler):
        server_version = "DingCheckin/1.0"
        protocol_version = "HTTP/1.1"

        # ---------------- 基础工具 ----------------
        def log_message(self, fmt, *args):
            app.logger.debug("web %s - %s", self.address_string(), fmt % args)

        def _send(self, body: bytes, content_type: str, status: int = 200,
                  extra_headers: Optional[dict] = None):
            self.send_response(status)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "no-store")
            for key, value in (extra_headers or {}).items():
                self.send_header(key, value)
            self.end_headers()
            if self.command != "HEAD":
                self.wfile.write(body)

        def _send_json(self, payload, status: int = 200):
            body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            self._send(body, "application/json; charset=utf-8", status)

        def _fail(self, status: int, message: str):
            self._send_json({"ok": False, "message": message}, status)

        def _query(self) -> dict:
            parsed = urlparse(self.path)
            flat = {}
            for key, values in parse_qs(parsed.query).items():
                flat[key] = values[0] if values else ""
            return flat

        def _body_json(self) -> dict:
            try:
                length = int(self.headers.get("Content-Length") or 0)
            except ValueError:
                return {}
            if length <= 0:
                return {}
            raw = self.rfile.read(length)
            try:
                data = json.loads(raw.decode("utf-8"))
            except (ValueError, UnicodeDecodeError):
                return {}
            return data if isinstance(data, dict) else {}

        def _require_user(self) -> Optional[str]:
            user_id = app.current_user(self.headers.get("Cookie"))
            if not app.is_allowed(user_id):
                self._fail(401, "未登录或不在白名单内")
                return None
            return user_id

        # ---------------- 路由 ----------------
        def do_GET(self):
            path = urlparse(self.path).path.rstrip("/") or "/"
            try:
                if path == "/":
                    self._send(PAGE_HTML.encode("utf-8"), "text/html; charset=utf-8")
                elif path == "/api/auth":
                    self._handle_auth()
                elif path == "/api/me":
                    self._handle_me()
                elif path == "/api/places":
                    self._handle_places()
                elif path == "/api/records":
                    self._handle_records()
                elif path == "/api/export.xlsx":
                    self._handle_export()
                else:
                    self._fail(404, "接口不存在")
            except Exception as e:
                app.logger.error("处理 %s 失败: %s", self.path, e, exc_info=True)
                self._fail(500, f"服务内部错误: {e}")

        def do_HEAD(self):
            self.do_GET()

        def do_POST(self):
            path = urlparse(self.path).path.rstrip("/") or "/"
            try:
                if path == "/api/login":
                    self._handle_login()
                else:
                    self._fail(404, "接口不存在")
            except Exception as e:
                app.logger.error("处理 %s 失败: %s", self.path, e, exc_info=True)
                self._fail(500, f"服务内部错误: {e}")

        # ---------------- 各接口实现 ----------------
        def _handle_auth(self):
            query = self._query()
            self._send_json({
                "ok": True,
                "mode": app.cfg.auth_mode,
                # 首页地址里可以用 $CORPID$ 占位，钉钉会替换成真实 corpId
                "corpId": query.get("corpid") or app.cfg.corp_id,
                "clientId": app.api.client_id,
                "agentId": app.cfg.agent_id,
            })

        def _handle_me(self):
            user_id = app.current_user(self.headers.get("Cookie"))
            allowed = app.is_allowed(user_id)
            name = ""
            if allowed and user_id and user_id != "(debug)":
                try:
                    name = app.name_of_factory([user_id])(user_id)
                except Exception:
                    name = ""
            self._send_json({
                "ok": True,
                "authenticated": allowed,
                "user_id": user_id if allowed else None,
                "name": name if name != user_id else "",
                "auth_mode": app.cfg.auth_mode,
            })

        def _handle_login(self):
            if app.cfg.auth_mode != "dingtalk":
                self._send_json({"ok": True, "user_id": "(debug)"})
                return

            body = self._body_json()
            code = str(body.get("code") or "").strip()
            if not code:
                self._fail(400, "缺少免登授权码")
                return

            try:
                info = app.api.get_userinfo_by_code(code)
            except Exception as e:
                app.logger.error("免登失败: %s", e)
                self._fail(401, f"免登失败: {e}")
                return

            user_id = info.get("userid")
            if not user_id:
                self._fail(401, "免登未返回 userId")
                return

            name = info.get("name") or ""
            if name:
                try:
                    app.user_store.upsert_users([{"user_id": user_id, "name": name}])
                except Exception as e:
                    app.logger.warning("写入员工姓名缓存失败: %s", e)

            allowed = user_id in app.cfg.admin_user_ids
            app.logger.info(
                "免登成功 | userId=%s | 姓名=%s | 白名单=%s", user_id, name, allowed
            )
            if not allowed:
                self._send_json({
                    "ok": False, "allowed": False, "user_id": user_id, "name": name,
                    "message": "你不在查看白名单内，请联系管理员",
                }, 403)
                return

            cookie = app.make_session(user_id)
            body_bytes = json.dumps(
                {"ok": True, "allowed": True, "user_id": user_id, "name": name},
                ensure_ascii=False,
            ).encode("utf-8")
            self._send(
                body_bytes, "application/json; charset=utf-8", 200,
                {"Set-Cookie": (
                    f"{SESSION_COOKIE}={cookie}; Path=/; Max-Age={app.cfg.session_ttl}; "
                    "HttpOnly; SameSite=Lax"
                )},
            )

        def _handle_places(self):
            if self._require_user() is None:
                return
            try:
                places = app.checkin_store.distinct_places()
            except Exception as e:
                self._fail(503, f"读取数据库失败（MongoDB 是否已启动？）: {e}")
                return
            self._send_json({"ok": True, "places": places})

        def _handle_records(self):
            if self._require_user() is None:
                return
            query = self._query()
            try:
                result = app.query(query.get("from"), query.get("to"), query.get("place"))
            except Exception as e:
                app.logger.error("查询失败: %s", e, exc_info=True)
                self._fail(503, f"查询失败（MongoDB 是否已启动？）: {e}")
                return
            result["ok"] = True
            self._send_json(result)

        def _handle_export(self):
            if self._require_user() is None:
                return
            query = self._query()
            try:
                result = app.query(query.get("from"), query.get("to"), query.get("place"))
            except Exception as e:
                app.logger.error("导出查询失败: %s", e, exc_info=True)
                self._fail(503, f"导出失败（MongoDB 是否已启动？）: {e}")
                return

            user_ids = sorted({
                row["user_id"] for row in result["detail"] if row.get("user_id")
            })
            name_of = app.name_of_factory(user_ids)
            exported_at = datetime.now(app.tz).strftime("%Y-%m-%d %H:%M:%S")
            filename = f"checkin_{result['from']}_{result['to']}.xlsx"
            chinese_name = f"外勤打卡记录_{result['from']}_{result['to']}.xlsx"

            summary_rows, detail_rows = export_row_groups(result)
            payload = build_workbook(
                summary_rows, detail_rows,
                {
                    "range_text": f"{result['from']} ~ {result['to']}",
                    "place": result["place"] or "全部",
                    "exported_at": exported_at,
                },
            )

            app.logger.info(
                "导出 Excel | 范围=%s~%s | 地点=%s | 汇总=%d 行 | 明细=%d 行 | %d 字节",
                result["from"], result["to"], result["place"] or "全部",
                len(summary_rows), len(detail_rows), len(payload),
            )
            self._send(
                payload,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                200,
                {"Content-Disposition": (
                    f'attachment; filename="{filename}"; '
                    f"filename*=UTF-8''{quote(chinese_name)}"
                )},
            )

    return Handler


# ======================================================================
# 页面
# ======================================================================
PAGE_HTML = r"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, maximum-scale=1, user-scalable=no">
<title>外勤打卡记录</title>
<script src="https://g.alicdn.com/dingding/dingtalk-jsapi/3.0.25/dingtalk.open.js"></script>
<style>
  * { box-sizing: border-box; -webkit-tap-highlight-color: transparent; }
  body { margin: 0; font-family: -apple-system, BlinkMacSystemFont, "PingFang SC", "Microsoft YaHei", sans-serif;
         background: #f5f6f7; color: #1f2329; font-size: 14px; }
  header { background: #1677ff; color: #fff; padding: 12px 16px; position: sticky; top: 0; z-index: 10; }
  header h1 { margin: 0; font-size: 16px; font-weight: 600; }
  header .who { font-size: 12px; opacity: .9; margin-top: 4px; }
  .wrap { padding: 12px; }
  .card { background: #fff; border-radius: 8px; padding: 12px; margin-bottom: 12px; box-shadow: 0 1px 3px rgba(0,0,0,.06); }
  .filters { display: flex; flex-wrap: wrap; gap: 8px; align-items: flex-end; }
  .field { display: flex; flex-direction: column; gap: 4px; }
  .field label { font-size: 12px; color: #646a73; }
  input[type=date], select { padding: 7px 8px; border: 1px solid #d9d9d9; border-radius: 6px; font-size: 14px;
                             background: #fff; color: #1f2329; min-width: 130px; }
  button { padding: 8px 14px; border: 0; border-radius: 6px; background: #1677ff; color: #fff;
           font-size: 14px; cursor: pointer; }
  button.ghost { background: #fff; color: #1677ff; border: 1px solid #1677ff; }
  button:disabled { opacity: .5; cursor: default; }
  .stat { color: #646a73; font-size: 13px; margin: 4px 0 8px; }
  h2 { font-size: 14px; margin: 0 0 8px; }
  .scroll { overflow-x: auto; -webkit-overflow-scrolling: touch; }
  table { border-collapse: collapse; width: 100%; white-space: nowrap; }
  th, td { border-bottom: 1px solid #eee; padding: 8px 10px; text-align: left; font-size: 13px; }
  th { background: #fafafa; color: #646a73; font-weight: 500; position: sticky; top: 0; }
  td.people { white-space: normal; min-width: 180px; }
  .empty { color: #8f959e; padding: 16px 0; text-align: center; }
  .msg { padding: 10px 12px; border-radius: 6px; margin-bottom: 10px; font-size: 13px; }
  .msg.err { background: #fff2f0; color: #cf1322; border: 1px solid #ffccc7; }
  .msg.info { background: #f0f7ff; color: #0958d9; border: 1px solid #bae0ff; }
  .row { display: flex; justify-content: space-between; align-items: center; }
  a.toggle { color: #1677ff; text-decoration: none; font-size: 13px; }
</style>
</head>
<body>
<header>
  <h1>外勤打卡记录</h1>
  <div class="who" id="who">正在识别身份…</div>
</header>

<div class="wrap">
  <div id="msg"></div>

  <div class="card">
    <div class="filters">
      <div class="field"><label>开始日期</label><input type="date" id="from"></div>
      <div class="field"><label>结束日期</label><input type="date" id="to"></div>
      <div class="field"><label>地点</label><select id="place"><option value="">全部</option></select></div>
      <button id="btnQuery">查询</button>
      <button id="btnExport" class="ghost">导出 Excel</button>
    </div>
  </div>

  <div class="card">
    <h2>汇总（地点 / 日期 — 到岗人）</h2>
    <div class="stat" id="stat"></div>
    <div class="scroll"><table id="summaryTable"></table></div>
    <div id="summaryEmpty" class="empty" style="display:none">该条件下没有记录</div>
  </div>

  <div class="card">
    <div class="row">
      <h2 style="margin:0">明细</h2>
      <a class="toggle" href="#" id="toggleDetail">展开</a>
    </div>
    <div id="detailBox" style="display:none; margin-top:8px">
      <div class="scroll"><table id="detailTable"></table></div>
    </div>
  </div>
</div>

<script>
(function () {
  var state = { authenticated: false };

  function el(id) { return document.getElementById(id); }
  function showMsg(text, kind) {
    var box = el('msg');
    box.innerHTML = '';
    if (!text) return;
    var div = document.createElement('div');
    div.className = 'msg ' + (kind || 'info');
    div.textContent = text;
    box.appendChild(div);
  }
  function fmtDate(d) {
    var m = ('0' + (d.getMonth() + 1)).slice(-2), day = ('0' + d.getDate()).slice(-2);
    return d.getFullYear() + '-' + m + '-' + day;
  }
  function getJSON(url) {
    return fetch(url, { credentials: 'same-origin' }).then(function (r) {
      return r.json().catch(function () { return {}; }).then(function (data) {
        if (!r.ok) throw new Error(data.message || ('HTTP ' + r.status));
        return data;
      });
    });
  }
  function postJSON(url, body) {
    return fetch(url, {
      method: 'POST', credentials: 'same-origin',
      headers: { 'Content-Type': 'application/json' },
      body: JSON.stringify(body)
    }).then(function (r) {
      return r.json().catch(function () { return {}; }).then(function (data) {
        if (!r.ok) { var e = new Error(data.message || ('HTTP ' + r.status)); e.data = data; throw e; }
        return data;
      });
    });
  }

  // ---- 免登 ----
  function requestAuthCode(corpId, clientId) {
    return new Promise(function (resolve, reject) {
      var dd = window.dd;
      if (!dd) { reject(new Error('请在钉钉客户端内打开本页面')); return; }
      var fail = function (err) { reject(new Error('获取免登授权码失败: ' + JSON.stringify(err))); };
      if (typeof dd.requestAuthCode === 'function') {
        dd.requestAuthCode({ corpId: corpId, clientId: clientId, onSuccess: function (r) { resolve(r.code); }, onFail: fail });
      } else if (dd.runtime && dd.runtime.permission && dd.runtime.permission.requestAuthCode) {
        dd.runtime.permission.requestAuthCode({ corpId: corpId, onSuccess: function (r) { resolve(r.code); }, onFail: fail });
      } else {
        reject(new Error('当前钉钉版本不支持免登，请升级钉钉客户端'));
      }
    });
  }

  function login() {
    return getJSON('/api/auth').then(function (auth) {
      if (auth.mode === 'none') return { ok: true };
      return requestAuthCode(auth.corpId, auth.clientId).then(function (code) {
        return postJSON('/api/login', { code: code });
      });
    });
  }

  // ---- 渲染 ----
  function renderTable(table, headers, rows, cellClass) {
    table.innerHTML = '';
    if (!rows.length) return;
    var thead = document.createElement('thead'), tr = document.createElement('tr');
    headers.forEach(function (h) { var th = document.createElement('th'); th.textContent = h; tr.appendChild(th); });
    thead.appendChild(tr); table.appendChild(thead);
    var tbody = document.createElement('tbody');
    rows.forEach(function (cells) {
      var row = document.createElement('tr');
      cells.forEach(function (value, index) {
        var td = document.createElement('td');
        td.textContent = (value === null || value === undefined) ? '' : String(value);
        if (cellClass && cellClass[index]) td.className = cellClass[index];
        row.appendChild(td);
      });
      tbody.appendChild(row);
    });
    table.appendChild(tbody);
  }

  function render(data) {
    el('stat').textContent = '共 ' + data.summary.length + ' 条地点/日期记录，'
      + data.record_total + ' 次打卡，' + data.user_total + ' 人';
    var rows = data.summary.map(function (row) {
      return [row.place, row.date, row.first_time, row.last_time,
              row.user_count, (row.users || []).join('、')];
    });
    renderTable(el('summaryTable'), ['地点', '日期', '首次打卡', '末次打卡', '到岗人数', '到岗人'],
                rows, { 5: 'people' });
    el('summaryEmpty').style.display = rows.length ? 'none' : 'block';

    var detailRows = data.detail.map(function (row) {
      return [row.name, row.place, row.date, row.time, row.distance,
              row.address, row.remark];
    });
    renderTable(el('detailTable'), ['姓名', '地点', '日期', '时间', '距离(米)', '打卡地址', '外勤备注'],
                detailRows);
  }

  function loadPlaces() {
    return getJSON('/api/places').then(function (data) {
      var select = el('place');
      (data.places || []).forEach(function (name) {
        var option = document.createElement('option');
        option.value = name; option.textContent = name;
        select.appendChild(option);
      });
    }).catch(function () { /* 地点列表失败不阻塞查询 */ });
  }

  function query() {
    var params = new URLSearchParams({
      from: el('from').value, to: el('to').value, place: el('place').value
    });
    showMsg('');
    return getJSON('/api/records?' + params.toString()).then(render)
      .catch(function (e) { showMsg(e.message, 'err'); });
  }

  function boot() {
    var today = new Date();
    el('to').value = fmtDate(today);
    var start = new Date(today.getTime() - 6 * 86400000);
    el('from').value = fmtDate(start);

    getJSON('/api/me').then(function (me) {
      if (me.authenticated) {
        state.authenticated = true;
        el('who').textContent = me.name ? (me.name + '（' + me.user_id + '）') : (me.user_id || '');
        return loadPlaces().then(query);
      }
      el('who').textContent = '正在免登…';
      return login().then(function () {
        return getJSON('/api/me').then(function (me2) {
          state.authenticated = !!me2.authenticated;
          el('who').textContent = me2.name || me2.user_id || '';
          return loadPlaces().then(query);
        });
      });
    }).catch(function (e) {
      el('who').textContent = '未登录';
      showMsg(e.message, 'err');
    });
  }

  el('btnQuery').addEventListener('click', query);
  el('btnExport').addEventListener('click', function () {
    var params = new URLSearchParams({
      from: el('from').value, to: el('to').value, place: el('place').value
    });
    window.location.href = '/api/export.xlsx?' + params.toString();
  });
  el('toggleDetail').addEventListener('click', function (event) {
    event.preventDefault();
    var box = el('detailBox');
    var open = box.style.display !== 'none';
    box.style.display = open ? 'none' : 'block';
    this.textContent = open ? '展开' : '收起';
  });

  boot();
})();
</script>
</body>
</html>
"""

#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
外勤打卡的查询聚合与 Excel 报表。

与触发入口无关（机器人 / 以后要加别的入口都能复用）：
  - resolve_range          把 YYYY-MM-DD 起止日期换算成 UTC 时间区间
  - summarize              按「地点 + 日期」聚合 -> 地点/日期/首次/末次/到岗人
  - detail_rows            逐条打卡明细
  - build_workbook         汇总 + 明细 两个 sheet 的 xlsx 字节流
  - format_summary_markdown  机器人聊天里展示用的文本
"""

import io
from datetime import datetime, timedelta
from datetime import time as dtime
from datetime import timezone
from typing import Callable, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_RANGE_DAYS = 7
DEFAULT_QUERY_LIMIT = 20000

SUMMARY_HEADERS = ["地点", "日期", "首次打卡", "末次打卡", "到岗人数", "到岗人"]
DETAIL_HEADERS = ["姓名", "userId", "地点", "打卡日期", "打卡时间",
                  "距离(米)", "打卡地址", "外勤备注", "打卡结果", "数据来源"]


# ======================================================================
# 时间区间
# ======================================================================
def resolve_range(from_text: Optional[str], to_text: Optional[str],
                  tz: timezone, default_days: int = DEFAULT_RANGE_DAYS):
    """把 YYYY-MM-DD 的起止日期换算成 UTC 的 [start, end] 闭区间。

    :return: (起始日期, 结束日期, start_utc, end_utc)
    """
    today = datetime.now(tz).date()
    try:
        to_date = datetime.strptime((to_text or "").strip(), "%Y-%m-%d").date()
    except (TypeError, ValueError):
        to_date = today
    try:
        from_date = datetime.strptime((from_text or "").strip(), "%Y-%m-%d").date()
    except (TypeError, ValueError):
        from_date = to_date - timedelta(days=max(default_days, 1) - 1)

    if from_date > to_date:
        from_date, to_date = to_date, from_date

    start_local = datetime.combine(from_date, dtime.min, tzinfo=tz)
    end_local = datetime.combine(to_date, dtime.max, tzinfo=tz)
    return (from_date, to_date,
            start_local.astimezone(timezone.utc),
            end_local.astimezone(timezone.utc))


def to_local(moment, tz: timezone) -> Optional[datetime]:
    """把库里的时间转成本地时间；naive 时间按 UTC 处理。"""
    if not isinstance(moment, datetime):
        return None
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=timezone.utc)
    return moment.astimezone(tz)


# ======================================================================
# 聚合
# ======================================================================
def summarize(records: List[dict], tz: timezone) -> List[dict]:
    """按「地点 + 日期」聚合，得到 地点/日期/首次/末次/人数/打卡次数/到岗人。"""
    buckets: dict = {}
    for record in records:
        local = to_local(record.get("user_check_time"), tz)
        if local is None:
            continue
        place = record.get("matched_place") or "(未知地点)"
        day = local.strftime("%Y-%m-%d")
        bucket = buckets.setdefault((place, day), {"users": set(), "times": [], "count": 0})
        bucket["count"] += 1
        bucket["times"].append(local)
        user_id = record.get("user_id")
        if user_id:
            bucket["users"].add(user_id)

    rows = []
    for place, day in sorted(buckets, key=lambda item: (item[1], item[0])):
        bucket = buckets[(place, day)]
        rows.append({
            "place": place,
            "date": day,
            "first_time": min(bucket["times"]).strftime("%H:%M:%S"),
            "last_time": max(bucket["times"]).strftime("%H:%M:%S"),
            "user_ids": sorted(bucket["users"]),
            "user_count": len(bucket["users"]),
            "record_count": bucket["count"],
        })
    return rows


def detail_rows(records: List[dict], tz: timezone, name_of: Callable[[str], str]) -> List[dict]:
    """明细行（按打卡时间正序）。"""
    prepared = []
    for record in records:
        local = to_local(record.get("user_check_time"), tz)
        if local is None:
            continue
        prepared.append((local, record))
    prepared.sort(key=lambda item: item[0])

    rows = []
    for local, record in prepared:
        user_id = record.get("user_id") or ""
        rows.append({
            "name": name_of(user_id),
            "user_id": user_id,
            "place": record.get("matched_place") or "",
            "date": local.strftime("%Y-%m-%d"),
            "time": local.strftime("%H:%M:%S"),
            "distance": record.get("distance_meters"),
            "address": record.get("user_address") or "",
            "remark": record.get("outside_remark") or "",
            "time_result": record.get("time_result") or "",
            "source_type": record.get("source_type") or "",
        })
    return rows


def build_result(checkin_store, user_store, tz: timezone,
                 from_text: Optional[str] = None, to_text: Optional[str] = None,
                 place: Optional[str] = None,
                 default_days: int = DEFAULT_RANGE_DAYS,
                 limit: int = DEFAULT_QUERY_LIMIT) -> dict:
    """查询 + 聚合，返回给展示层用的结果结构。

    MongoDB 不可用时会让异常抛出，由调用方给出友好提示。
    """
    from_date, to_date, start_utc, end_utc = resolve_range(
        from_text, to_text, tz, default_days=default_days
    )
    records = checkin_store.find_records(
        start_utc, end_utc, place=place or None, limit=limit
    )

    user_ids = sorted({r.get("user_id") for r in records if r.get("user_id")})
    mapping = {}
    if user_ids and user_store is not None:
        try:
            mapping = user_store.name_map(user_ids) or {}
        except Exception:
            mapping = {}
    name_of = lambda uid: mapping.get(uid) or (uid or "-")   # noqa: E731

    return {
        "from": from_date.strftime("%Y-%m-%d"),
        "to": to_date.strftime("%Y-%m-%d"),
        "place": place or "",
        "record_total": len(records),
        "user_total": len(user_ids),
        "summary": [
            {
                "place": row["place"],
                "date": row["date"],
                "first_time": row["first_time"],
                "last_time": row["last_time"],
                "user_count": row["user_count"],
                "record_count": row["record_count"],
                "users": [name_of(uid) for uid in row["user_ids"]],
            }
            for row in summarize(records, tz)
        ],
        "detail": detail_rows(records, tz, name_of),
    }


# ======================================================================
# 聊天里的文本展示
# ======================================================================
def format_summary_markdown(result: dict, max_rows: int = 30,
                            max_names: int = 60) -> str:
    """把查询结果渲染成机器人回复的 markdown。

    每条形如：  **地点** ｜ 日期 ｜ 首次~末次 ｜ N人
                到岗人：张三、李四
    """
    header = [
        "### 外勤打卡记录",
        f"**范围**：{result['from']} ~ {result['to']}"
        f"　**地点**：{result['place'] or '全部'}",
        f"**合计**：{len(result['summary'])} 条（地点/日期）｜"
        f"{result['record_total']} 次打卡｜{result['user_total']} 人",
    ]

    rows = result["summary"]
    if not rows:
        return "  \n".join(header + ["", "该条件下没有外勤打卡记录。"])

    blocks = []
    for row in rows[:max_rows]:
        names = row["users"]
        shown = "、".join(names[:max_names])
        if len(names) > max_names:
            shown += f" 等 {len(names)} 人"
        blocks.append(
            f"**{row['place']}** ｜ {row['date']} ｜ "
            f"{row['first_time']}~{row['last_time']} ｜ {row['user_count']} 人"
            f"  \n到岗人：{shown or '-'}"
        )

    if len(rows) > max_rows:
        blocks.append(f"…… 还有 {len(rows) - max_rows} 条未显示，请用「导出」获取完整 Excel。")

    return "  \n".join(header + [""] + blocks)


# ======================================================================
# Excel
# ======================================================================
def export_row_groups(result: dict) -> Tuple[List[dict], List[dict]]:
    """把查询结果转成写 Excel 需要的两组行（到岗人已经是姓名）。"""
    summary = [
        {
            "place": row["place"],
            "date": row["date"],
            "first_time": row["first_time"],
            "last_time": row["last_time"],
            "user_count": row["user_count"],
            "people": "、".join(row.get("users") or []),
        }
        for row in result["summary"]
    ]
    return summary, result["detail"]


def build_workbook(summary_rows: List[dict], detail_rows: List[dict],
                   meta: dict) -> bytes:
    """生成 xlsx（汇总 + 明细 两个 sheet），返回字节流。

    :param summary_rows: [{place, date, first_time, last_time, user_count, people}]
    :param detail_rows:  明细行（字段顺序见 DETAIL_HEADERS）
    :param meta:         查询条件，写进汇总表尾部便于回溯口径
    """
    workbook = Workbook()

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")

    def write_sheet(sheet, headers, rows, widths):
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        if rows:
            sheet.auto_filter.ref = sheet.dimensions
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width

    sheet = workbook.active
    sheet.title = "汇总"
    write_sheet(
        sheet, SUMMARY_HEADERS,
        [[row["place"], row["date"], row["first_time"], row["last_time"],
          row["user_count"], row.get("people", "")] for row in summary_rows],
        [24, 12, 10, 10, 10, 60],
    )

    detail_sheet = workbook.create_sheet("明细")
    write_sheet(
        detail_sheet, DETAIL_HEADERS,
        [[row["name"], row["user_id"], row["place"], row["date"], row["time"],
          row["distance"], row["address"], row["remark"],
          row["time_result"], row["source_type"]] for row in detail_rows],
        [12, 24, 24, 12, 10, 10, 32, 20, 10, 10],
    )

    sheet.append([])
    sheet.append(["查询条件", str(meta.get("range_text") or ""),
                  "地点", str(meta.get("place") or "全部"),
                  "导出时间", str(meta.get("exported_at") or "")])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
